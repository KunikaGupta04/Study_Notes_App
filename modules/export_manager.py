from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF
import io
import re


def strip_markdown(text):
    """Remove markdown syntax for plain-text export."""
    text = re.sub(r'```.*?```', '', text, flags=re.DOTALL)   # code blocks first
    text = re.sub(r'`[^`]+`', '', text)                       # inline code
    text = re.sub(r'#{1,6}\s+', '', text)                     # headings
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)              # bold
    text = re.sub(r'\*(.*?)\*', r'\1', text)                  # italic
    text = re.sub(r'\|.*?\|', '', text)                       # tables
    text = re.sub(r'^[-*+]\s+', '• ', text, flags=re.MULTILINE)
    text = re.sub(r'\n{3,}', '\n\n', text)                    # collapse blank lines
    return text.strip()

def generate_word(content):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Pt(60)
    section.bottom_margin = Pt(60)
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)

    # Title
    title = doc.add_paragraph()
    title.alignment = 1
    run = title.add_run('AI Smart Notes')
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x63, 0x66, 0xF1)

    doc.add_paragraph()

    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            doc.add_paragraph()
            i += 1
            continue

        if line.strip().startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                i += 1
            i += 1
            continue

        if line.strip().startswith('|'):
            i += 1
            continue

        if line.startswith('# ') and not line.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(line[2:].strip())
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x63, 0x66, 0xF1)

        elif line.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(line[3:].strip())
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)

        elif line.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(line[4:].strip())
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x55, 0x48, 0xE5)

        elif line.startswith(('- ', '* ', '+ ')):
            text = line[2:].strip()
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run('• ' + text)
            run.font.size = Pt(11)

        else:
            clean = re.sub(r'\*\*(.*?)\*\*', r'\1', line.strip())
            clean = re.sub(r'\*(.*?)\*', r'\1', clean)
            clean = re.sub(r'`(.*?)`', r'\1', clean)
            if clean:
                p = doc.add_paragraph()
                run = p.add_run(clean)
                run.font.size = Pt(11)

        i += 1

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
def generate_pdf(content):
    """Generate a PDF from raw markdown notes. Compatible with fpdf2."""
    import re, io
    from fpdf import FPDF

    # ── strip markdown + emojis (latin-1 can't encode emojis) ──
    def clean(text):
        text = re.sub(r'```.*?```', '', text, flags=re.DOTALL)   # code/mermaid blocks
        text = re.sub(r'`[^`]+`', '', text)                       # inline code
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)              # bold
        text = re.sub(r'\*(.*?)\*',     r'\1', text)              # italic
        text = re.sub(r'\|.*?\|', '', text)                       # tables
        text = re.sub(r'^[-*+]\s+', '• ', text, flags=re.MULTILINE)
        text = re.sub(r'\n{3,}', '\n\n', text)
        # Remove emojis and non-latin-1 characters
        text = text.encode('latin-1', 'ignore').decode('latin-1')
        return text.strip()

    PAGE_W   = 210          # A4 mm
    MARGIN   = 20
    CONTENT_W = PAGE_W - 2 * MARGIN   # 170 mm — always positive

    pdf = FPDF()
    pdf.set_margins(MARGIN, MARGIN, MARGIN)
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=MARGIN)

    # Title
    pdf.set_font("Helvetica", 'B', 18)
    pdf.cell(CONTENT_W, 12, "AI Smart Notes", align='C', new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    lines = clean(content).split('\n')

    for line in lines:
        line = line.strip()
        if not line:
            pdf.ln(4)
            continue

        if line.startswith('# ') or (line.startswith('## ') and not line.startswith('### ')):
            heading = line.lstrip('#').strip()
            pdf.set_font("Helvetica", 'B', 14)
            pdf.ln(3)
            pdf.multi_cell(CONTENT_W, 9, heading, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)

        elif line.startswith('### '):
            heading = line[4:].strip()
            pdf.set_font("Helvetica", 'B', 12)
            pdf.multi_cell(CONTENT_W, 8, heading, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)

        elif line.startswith('• '):
            pdf.set_font("Helvetica", size=11)
            # fixed indent: reduce content width by indent amount, move x right
            INDENT = 8
            pdf.set_x(MARGIN + INDENT)
            pdf.multi_cell(CONTENT_W - INDENT, 7, line, new_x="LMARGIN", new_y="NEXT")

        else:
            pdf.set_font("Helvetica", size=11)
            pdf.multi_cell(CONTENT_W, 7, line, new_x="LMARGIN", new_y="NEXT")

    pdf_bytes = pdf.output()
    if isinstance(pdf_bytes, str):
        pdf_bytes = pdf_bytes.encode('latin-1')
    return io.BytesIO(bytes(pdf_bytes))